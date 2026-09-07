"""Workbook / sheet inventory for every quotation source file."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .extract import SOURCE_DIR, detect_header, source_files, _as_text
from .classify import document_type


def hash_inventory(directory: Path = SOURCE_DIR) -> list[dict]:
    unique, skipped = source_files(directory)
    rows = []
    for rec in unique:
        rows.append(
            {
                "FILE": rec["name"],
                "PATH": rec["path"],
                "SHA256": rec["sha256"],
                "STATUS": "UNIQUE",
                "CANONICAL_FILE": rec["name"],
            }
        )
    for rec in skipped:
        rows.append(
            {
                "FILE": rec["name"],
                "PATH": rec["path"],
                "SHA256": rec["sha256"],
                "STATUS": "EXACT_DUPLICATE_FILE_SKIPPED",
                "CANONICAL_FILE": rec["duplicate_of"],
            }
        )
    return rows


def sheet_inventory(directory: Path = SOURCE_DIR) -> list[dict]:
    unique, skipped = source_files(directory)
    skip_names = {s["name"] for s in skipped}
    all_recs = list(unique) + list(skipped)
    out = []
    for rec in all_recs:
        path = Path(rec["path"])
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for title in wb.sheetnames:
                hidden = False
                try:
                    hidden = wb[title].sheet_state != "visible"
                except Exception:
                    hidden = False
                ws = wb[title]
                rows = list(
                    ws.iter_rows(max_row=min(ws.max_row or 0, 400), values_only=True)
                )
                nonempty = sum(1 for r in rows if any(_as_text(c) for c in r))
                cols = max((len(r) for r in rows), default=0)
                preview = " ".join(
                    _as_text(c) for r in rows[:20] for c in r[:8] if _as_text(c)
                )[:400]
                header, mapping = detect_header(rows)
                doc = document_type(title, preview) if nonempty else "EMPTY"
                has_lines = bool(mapping.get("desc") is not None and nonempty > 2)
                if rec["name"] in skip_names:
                    status = "SKIP_DUPLICATE_FILE"
                elif nonempty == 0:
                    status = "EMPTY"
                    doc = "EMPTY"
                elif not has_lines and nonempty < 3:
                    status = "EMPTY"
                    doc = "EMPTY"
                elif not has_lines:
                    status = "NO_HEADER_REVIEW"
                else:
                    status = "PROCESS"
                out.append(
                    {
                        "FILE": rec["name"],
                        "SHA256": rec["sha256"],
                        "SHEET": title,
                        "VISIBLE": "HIDDEN" if hidden else "VISIBLE",
                        "ROWS": ws.max_row or 0,
                        "COLUMNS": cols,
                        "NONEMPTY_ROWS_SAMPLED": nonempty,
                        "DOCUMENT_TYPE": doc,
                        "HAS_PRODUCT_LINES": has_lines,
                        "HEADER_ROW": header if header is not None else "",
                        "STATUS": status,
                    }
                )
        finally:
            wb.close()
    return out
